"""读取 TikTok SKU 导出表。"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ASIN_RE = re.compile(r"^B0[A-Z0-9]{8}$", re.I)

# 需求关注字段（兼容表头）
REQUIRED_HINTS = ["平台SKU", "本地展示价"]


def _clean(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, str):
        return v.strip().lstrip("\t").strip()
    return v


def load_sku_rows(path: str | Path) -> list[dict[str, Any]]:
    path = Path(path)
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = None
    rows: list[tuple[Any, ...]] = []
    available_headers: dict[str, list[str]] = {}
    try:
        for sheet in wb.worksheets:
            sheet_rows = list(sheet.iter_rows(values_only=True))
            if not sheet_rows:
                available_headers[sheet.title] = []
                continue
            sheet_headers = [
                str(h).strip() if h is not None else f"col{i}"
                for i, h in enumerate(sheet_rows[0])
            ]
            available_headers[sheet.title] = sheet_headers
            if all(required in sheet_headers for required in REQUIRED_HINTS):
                ws = sheet
                rows = sheet_rows
                break
    finally:
        wb.close()
    if ws is None:
        raise ValueError(
            "Excel 中没有找到同时包含 平台SKU/本地展示价 的工作表，"
            f"已检查: {available_headers}"
        )
    headers = [str(h).strip() if h is not None else f"col{i}" for i, h in enumerate(rows[0])]
    idx = {h: i for i, h in enumerate(headers)}

    def col(*names, default=None):
        for n in names:
            if n in idx:
                return idx[n]
        return default

    i_asin = col("平台SKU")
    i_price = col("本地展示价")
    if i_asin is None or i_price is None:
        raise ValueError(f"Excel 缺少必要列 平台SKU/本地展示价，当前表头: {headers}")

    i_sku = col("SKU ID")
    i_pid = col("产品ID")
    i_name = col("产品名称")
    i_spec = col("规格")
    i_stock = col("库存")
    i_src_price = col("关联货源价格")
    i_src_link = col("关联货源链接")

    out: list[dict[str, Any]] = []
    for n, r in enumerate(rows[1:], start=2):
        if r is None or all(c is None or str(c).strip() == "" for c in r):
            continue
        asin = _clean(r[i_asin] if i_asin < len(r) else None)
        asin = str(asin) if asin is not None else ""
        local = r[i_price] if i_price < len(r) else None
        if isinstance(local, str):
            local_s = local.replace("$", "").replace(",", "").strip()
            try:
                local = float(local_s)
            except Exception:
                local = None

        item = {
            "row_num": n,
            "sku_id": _clean(r[i_sku]) if i_sku is not None else None,
            "product_id": _clean(r[i_pid]) if i_pid is not None else None,
            "product_name": _clean(r[i_name]) if i_name is not None else None,
            "spec": _clean(r[i_spec]) if i_spec is not None else None,
            "asin": asin,
            "local_price": local,
            "stock": r[i_stock] if i_stock is not None else None,
            "source_price": r[i_src_price] if i_src_price is not None else None,
            "source_link": _clean(r[i_src_link]) if i_src_link is not None else None,
            "asin_valid": bool(asin and ASIN_RE.match(asin)),
            "amazon_link": f"https://www.amazon.com/dp/{asin}" if asin else None,
        }
        out.append(item)
    return out

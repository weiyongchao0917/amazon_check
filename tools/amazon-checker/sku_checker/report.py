"""写出检查报告 Excel。"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


def needs_retry(row: dict[str, Any]) -> bool:
    """仅标记本次未成功取得 Amazon 数据、值得再次联网检查的条目。"""
    return (
        row.get("sync_status") == "同步失败已跳过"
        or row.get("overall") == "接口数据不可用"
    )


def write_report(
    rows: list[dict[str, Any]],
    out_path: str | Path,
    *,
    meta: dict[str, Any] | None = None,
) -> Path:
    meta = meta or {}
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cell_font = Font(name="Arial", size=10)
    title_font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    thin = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )
    header_fill = PatternFill("solid", fgColor="1F4E79")
    title_fill = PatternFill("solid", fgColor="1F4E79")
    info_fill = PatternFill("solid", fgColor="DDEBF7")
    ok_fill = PatternFill("solid", fgColor="C6EFCE")
    warn_fill = PatternFill("solid", fgColor="FFEB9C")
    bad_fill = PatternFill("solid", fgColor="FFC7CE")

    def colorize(cell, text: str) -> None:
        if text in ("正常", "价格正常", "配送正常", "无需处理"):
            cell.fill = ok_fill
        elif text in ("价格需检查", "配送需检查", "价格和配送均需检查", "配送需人工确认"):
            cell.fill = warn_fill
        elif text:
            cell.fill = bad_fill

    headers = [
        "序号",
        "SKU ID",
        "产品ID",
        "产品名称",
        "TikTok规格",
        "Amazon主规格",
        "Amazon规格维度",
        "规格检查结果",
        "规格判定依据",
        "平台SKU",
        "Amazon链接",
        "本地展示价",
        "Amazon当前价格",
        "建议本地价",
        "价格差异金额",
        "价格差异百分比",
        "价格检查结果",
        "Amazon库存状态/数量",
        "库存建议",
        "货源/报价状态",
        "状态判定依据",
        "Amazon预计送达日期",
        "配送方案类型",
        "配送提取来源",
        "配送方案原文",
        "配送诊断信息",
        "允许最晚送达日期",
        "配送检查结果",
        "总状态",
        "人工处理建议",
        "目标邮编",
        "页面邮编显示",
        "邮编确认状态",
        "是否有加购",
        "代理",
        "同步状态",
        "尝试次数",
        "接口来源",
        "错误信息",
        "检查时间",
    ]

    ws = wb.active
    ws.title = "检查报告"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws["A1"] = "SKU 检查报告"
    ws["A1"].font = title_font
    ws["A1"].fill = title_fill
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws["A2"] = (
        f"检查时间: {meta.get('checked_at','')} | 邮编: {meta.get('zip_code','')} | "
        f"倍率: {meta.get('price_multiplier','')} | 阈值: {meta.get('price_threshold','')} | "
        f"库存预警阈值: {meta.get('stock_warning_threshold','')} | "
        f"允许最晚送达日期: {meta.get('latest_allowed_date','')} | "
        f"代理: {meta.get('proxy_mode','none')} | "
        f"规则: Amazon最快配送日期不晚于用户指定截止日期 / 日期范围取第一天 / Amazon当前可见购买价优先"
    )
    ws["A2"].font = Font(name="Arial", size=9, italic=True)
    ws["A2"].fill = info_fill
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 40

    for c, h in enumerate(headers, 1):
        cell = ws.cell(3, c, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin
        cell.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")

    for i, r in enumerate(rows, 1):
        diff_pct = r.get("diff_pct")
        vals = [
            i,
            r.get("sku_id"),
            r.get("product_id"),
            r.get("product_name"),
            r.get("spec"),
            r.get("amazon_spec"),
            r.get("amazon_spec_dimension"),
            r.get("spec_flag"),
            r.get("spec_reason"),
            r.get("asin"),
            r.get("amazon_link"),
            r.get("local_price"),
            r.get("amazon_price"),
            r.get("suggested_local"),
            r.get("diff_amount"),
            (diff_pct / 100.0) if isinstance(diff_pct, (int, float)) else None,
            r.get("price_flag"),
            r.get("stock_status"),
            r.get("stock_advice"),
            r.get("source_status"),
            r.get("status_reason"),
            r.get("delivery_time"),
            r.get("delivery_option_type"),
            r.get("delivery_extract_source"),
            r.get("delivery_text"),
            r.get("delivery_diagnostic"),
            r.get("allow_date"),
            r.get("delivery_flag"),
            r.get("overall"),
            r.get("advice"),
            r.get("target_zip") or meta.get("zip_code"),
            r.get("location_line"),
            r.get("zip_confirmed"),
            r.get("has_add_to_cart"),
            r.get("proxy_used"),
            r.get("sync_status"),
            r.get("attempt_count"),
            r.get("source") or "Amazon公开页",
            r.get("error"),
            r.get("checked_at"),
        ]
        row_i = 3 + i
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row_i, c, v)
            cell.font = cell_font
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            if c in (12, 13, 14, 15) and isinstance(v, (int, float)):
                cell.number_format = "0.00"
            if c == 16 and isinstance(v, float):
                cell.number_format = "0.00%"
        colorize(ws.cell(row_i, 8), str(r.get("spec_flag") or ""))
        colorize(ws.cell(row_i, 17), str(r.get("price_flag") or ""))
        colorize(ws.cell(row_i, 19), str(r.get("stock_advice") or ""))
        colorize(ws.cell(row_i, 28), str(r.get("delivery_flag") or ""))
        colorize(ws.cell(row_i, 29), str(r.get("overall") or ""))

    widths = [6, 18, 16, 32, 16, 16, 14, 14, 32, 14, 32, 12, 12, 12, 12, 12, 14, 14, 12, 20, 34, 18, 14, 18, 36, 48, 14, 12, 18, 30, 12, 18, 12, 10, 22, 18, 10, 14, 24, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A4"
    if rows:
        ws.auto_filter.ref = f"A3:{get_column_letter(len(headers))}{3+len(rows)}"

    # 异常汇总
    ws2 = wb.create_sheet("异常汇总")
    ws2["A1"] = "需要人工处理（总状态 ≠ 正常）"
    ws2["A1"].font = title_font
    ws2["A1"].fill = title_fill
    ws2.merge_cells("A1:U1")
    abn_h = [
        "序号",
        "SKU ID",
        "全球产品ID",
        "平台SKU",
        "产品名称",
        "TikTok规格",
        "Amazon主规格",
        "规格检查结果",
        "本地展示价",
        "Amazon当前价格",
        "建议本地价",
        "价格检查结果",
        "库存状态/数量",
        "库存建议",
        "货源/报价状态",
        "状态判定依据",
        "送达日期",
        "配送检查结果",
        "总状态",
        "处理建议",
        "Amazon链接",
    ]
    for c, h in enumerate(abn_h, 1):
        cell = ws2.cell(2, c, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin
    abn = [r for r in rows if r.get("overall") != "正常"]
    for i, r in enumerate(abn, 1):
        vals = [
            i,
            r.get("sku_id"),
            r.get("product_id"),
            r.get("asin"),
            r.get("product_name"),
            r.get("spec"),
            r.get("amazon_spec"),
            r.get("spec_flag"),
            r.get("local_price"),
            r.get("amazon_price"),
            r.get("suggested_local"),
            r.get("price_flag"),
            r.get("stock_status"),
            r.get("stock_advice"),
            r.get("source_status"),
            r.get("status_reason"),
            r.get("delivery_time"),
            r.get("delivery_flag"),
            r.get("overall"),
            r.get("advice"),
            r.get("amazon_link"),
        ]
        for c, v in enumerate(vals, 1):
            cell = ws2.cell(2 + i, c, v)
            cell.font = cell_font
            cell.border = thin
        colorize(ws2.cell(2 + i, 8), str(r.get("spec_flag") or ""))
        colorize(ws2.cell(2 + i, 12), str(r.get("price_flag") or ""))
        colorize(ws2.cell(2 + i, 14), str(r.get("stock_advice") or ""))
        colorize(ws2.cell(2 + i, 18), str(r.get("delivery_flag") or ""))
        colorize(ws2.cell(2 + i, 19), str(r.get("overall") or ""))
    for i, w in enumerate([6, 18, 18, 14, 32, 16, 16, 14, 12, 12, 12, 12, 14, 12, 20, 34, 16, 12, 18, 30, 32], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w

    # 仅列出因网络、代理或页面取数失败而需要重新联网检查的条目。
    ws_retry = wb.create_sheet("需要重新查")
    ws_retry["A1"] = "本次未取到完整数据，可在 GUI 点击“重试失败项”"
    ws_retry["A1"].font = title_font
    ws_retry["A1"].fill = title_fill
    ws_retry.merge_cells("A1:J1")
    retry_headers = [
        "序号", "SKU ID", "产品名称", "平台SKU", "同步状态",
        "总状态", "错误信息", "尝试次数", "代理", "Amazon链接",
    ]
    for c, h in enumerate(retry_headers, 1):
        cell = ws_retry.cell(2, c, h)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = thin
    retry_rows = [r for r in rows if needs_retry(r)]
    for i, r in enumerate(retry_rows, 1):
        vals = [
            i, r.get("sku_id"), r.get("product_name"), r.get("asin"),
            r.get("sync_status"), r.get("overall"), r.get("error"),
            r.get("attempt_count"), r.get("proxy_used"), r.get("amazon_link"),
        ]
        for c, v in enumerate(vals, 1):
            cell = ws_retry.cell(2 + i, c, v)
            cell.font = cell_font
            cell.border = thin
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        colorize(ws_retry.cell(2 + i, 6), str(r.get("overall") or ""))
    for i, w in enumerate([6, 18, 32, 14, 18, 18, 60, 10, 24, 32], 1):
        ws_retry.column_dimensions[get_column_letter(i)].width = w
    ws_retry.freeze_panes = "A3"
    if retry_rows:
        ws_retry.auto_filter.ref = f"A2:J{2 + len(retry_rows)}"

    # 运行信息
    ws3 = wb.create_sheet("运行信息")
    ws3["A1"] = "运行信息"
    ws3["A1"].font = title_font
    ws3["A1"].fill = title_fill
    ws3.merge_cells("A1:B1")
    stats = meta.get("stats") or {}
    info_rows = [
        ("检查时间", meta.get("checked_at")),
        ("输入文件", meta.get("input_excel")),
        ("输出文件", str(out_path)),
        ("SKU总数", len(rows)),
        ("代理模式", meta.get("proxy_mode")),
        ("邮编", meta.get("zip_code")),
        ("允许最晚送达日期", meta.get("latest_allowed_date")),
        ("库存预警阈值", meta.get("stock_warning_threshold")),
        ("价格成功率", stats.get("price_ok")),
        ("配送日期成功率", stats.get("delivery_ok")),
        ("目标邮编确认率", stats.get("zip_ok")),
        ("访问受限次数", stats.get("blocked")),
        ("错误数", stats.get("errors")),
        ("总状态分布", stats.get("overall_dist")),
        ("说明", "默认直连；可在 config 启用代理列表或提取 API。不做验证码自动破解。"),
    ]
    ws3["A3"] = "项目"
    ws3["B3"] = "内容"
    for c in (1, 2):
        ws3.cell(3, c).font = header_font
        ws3.cell(3, c).fill = header_fill
        ws3.cell(3, c).border = thin
    for i, (k, v) in enumerate(info_rows, 4):
        ws3.cell(i, 1, k).font = cell_font
        ws3.cell(i, 1).border = thin
        ws3.cell(i, 2, str(v) if v is not None else "").font = cell_font
        ws3.cell(i, 2).border = thin
        ws3.cell(i, 2).alignment = Alignment(wrap_text=True)
    ws3.column_dimensions["A"].width = 18
    ws3.column_dimensions["B"].width = 100

    wb.save(out_path)
    return out_path

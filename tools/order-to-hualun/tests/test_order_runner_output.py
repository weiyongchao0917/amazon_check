import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from sku_checker.order_runner import run_order_conversion


class OrderOutputTests(unittest.TestCase):
    def test_output_contains_only_current_input_rows(self):
        with TemporaryDirectory() as temp:
            root = Path(temp)
            input_path = root / "orders.xlsx"
            output_path = root / "output.xlsx"
            state_path = root / "state.json"

            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["订单编号", "平台SKU", "产品数量", "产品规格(原文)"])
            for index in range(5):
                sheet.append([f"ORDER-{index}", f"B0TEST{index}", 1, "Default"])
            workbook.save(input_path)
            workbook.close()

            config = {
                "recipient_prefix": "Claire-wyc",
                "phone": "6265225687",
                "address1": "10136 Stafford St",
                "address2": "",
                "city": "Rancho Cucamonga",
                "state": "CA",
                "zip_code": "91730",
                "template_path": str(Path(__file__).parents[1] / "templates" / "花轮模板.xlsx"),
                "state_file": str(state_path),
            }

            run_order_conversion(
                str(input_path),
                config,
                output_path=str(output_path),
                fetcher=lambda row: {"price": 9.99, "note": "", "retryable": False},
            )

            result = load_workbook(output_path, data_only=False)
            try:
                result_sheet = result[result.sheetnames[0]]
                populated_rows = [
                    row
                    for row in range(2, result_sheet.max_row + 1)
                    if any(result_sheet.cell(row, column).value not in (None, "") for column in range(1, 15))
                ]
                self.assertEqual(populated_rows, [2, 3, 4, 5, 6])
            finally:
                result.close()


if __name__ == "__main__":
    unittest.main()

"""Order Excel -> Hualun template orchestration."""
from __future__ import annotations

import json
import hashlib
import shutil
from datetime import datetime
from pathlib import Path
from threading import Event
from typing import Any, Callable

from openpyxl.cell.cell import MergedCell
from openpyxl import load_workbook

from sku_checker.order_amazon import fetch_order_item
from sku_checker.order_io import OutputRow, load_order_rows, make_output_rows
from sku_checker.paths import ORDER_TEMPLATE_PATH, RUNTIME_DIR, configure_curl_ca_bundle

STATE_VERSION = 3


def _input_fingerprint(input_path: Path) -> str:
    """Return a content fingerprint so replaced files at the same path re-run."""
    digest = hashlib.sha256()
    with input_path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _compatible_done(state: dict[str, Any], input_path: Path) -> dict[str, Any]:
    """Return checkpoint rows only when they belong to this exact input file."""
    if state.get("state_version") != STATE_VERSION:
        return {}
    if str(state.get("input_excel") or "") != str(input_path.resolve()):
        return {}
    try:
        if str(state.get("input_fingerprint") or "") != _input_fingerprint(input_path):
            return {}
    except OSError:
        return {}
    done = state.get("done")
    return done if isinstance(done, dict) else {}


def _starting_done(saved_done: dict[str, Any], *, retry_failed_only: bool) -> dict[str, Any]:
    """Fresh runs query every row; retry runs keep only completed successes."""
    if not retry_failed_only:
        return {}
    return {
        key: value
        for key, value in saved_done.items()
        if not value.get("retryable", True)
    }


def _state_path(config: dict[str, Any]) -> Path:
    return Path(config.get("state_file") or (RUNTIME_DIR / "order_state.json"))


def _load_state(path: Path) -> dict[str, Any]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {"done": {}}
    except (OSError, ValueError):
        return {"done": {}}


def _save_state(path: Path, state: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + ".tmp")
    temp.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(path)


def _key(row: OutputRow) -> str:
    return f"{row.source.row_num}:{row.sku}"


def _default_output(input_path: Path, output_dir: str | Path | None) -> Path:
    directory = Path(output_dir) if output_dir else input_path.parent
    return directory / f"花轮采购订单-{datetime.now():%Y%m%d_%H%M%S}.xlsx"


def _note_text(result: Any) -> str:
    value = result.get("note") if isinstance(result, dict) else getattr(result, "note", "")
    return str(value or "").strip()


def _result_value(result: Any, key: str, default: Any = None) -> Any:
    if isinstance(result, dict):
        return result.get(key, default)
    return getattr(result, key, default)


def _clear_template_data(worksheet: Any) -> None:
    """Clear old order values while preserving the template's formatting."""
    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=14,
    ):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            cell.value = None
            cell.hyperlink = None
            cell.comment = None


def run_order_conversion(
    input_excel: str,
    config: dict[str, Any],
    *,
    output_path: str | None = None,
    fetcher: Callable[[Any], Any] | None = None,
    client: Any | None = None,
    cancel_event: Event | None = None,
    progress_callback: Callable[[dict[str, Any]], None] | None = None,
    retry_failed_only: bool = False,
) -> Path:
    input_path = Path(input_excel)
    rows = load_order_rows(input_path)
    output_rows = make_output_rows(rows, str(config.get("recipient_prefix") or "Claire-wyc"))
    input_fingerprint = _input_fingerprint(input_path)
    state_path = _state_path(config)
    state = _load_state(state_path)
    saved_done = _compatible_done(state, input_path)
    done: dict[str, Any] = _starting_done(saved_done, retry_failed_only=retry_failed_only)
    if not saved_done and (
        state.get("state_version") != STATE_VERSION
        or state.get("input_excel") != str(input_path.resolve())
        or state.get("input_fingerprint") != input_fingerprint
    ):
        state = {"state_version": STATE_VERSION, "done": {}}
    if fetcher is None:
        if client is None:
            configure_curl_ca_bundle()
            from sku_checker.amazon_client import AmazonClient
            client = AmazonClient(zip_code=str(config.get("zip_code") or "91730"), max_retries=2, delay_range=(1.5, 3.5), proxy_pool=None)
        fetcher = lambda row: fetch_order_item(client, row.source)

    total = len(output_rows)
    for index, row in enumerate(output_rows, start=1):
        key = _key(row)
        if key in done:
            continue
        if cancel_event is not None and cancel_event.is_set():
            break
        if progress_callback:
            progress_callback({"event": "row_started", "current": index, "total": total, "sku": row.sku})
        try:
            result = fetcher(row)
            done[key] = {
                "price": _result_value(result, "price"),
                "note": _note_text(result),
                "retryable": bool(_result_value(result, "retryable", False)),
            }
        except Exception as exc:
            done[key] = {"price": None, "note": f"{type(exc).__name__}: {exc}", "retryable": True}
        state.update({"state_version": STATE_VERSION, "input_excel": str(input_path.resolve()), "input_fingerprint": input_fingerprint, "total": total, "done": done, "updated_at": datetime.now().isoformat()})
        _save_state(state_path, state)
        if progress_callback:
            progress_callback({"event": "row_completed", "current": index, "total": total, "sku": row.sku, "note": done[key]["note"]})

    template_path = Path(config.get("template_path") or ORDER_TEMPLATE_PATH)
    if not template_path.is_file():
        raise FileNotFoundError(f"内置花轮模板不存在: {template_path}")
    target = Path(output_path) if output_path else _default_output(input_path, config.get("output_dir"))
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, target)
    wb = load_workbook(target)
    try:
        ws = wb[wb.sheetnames[0]]
        _clear_template_data(ws)
        store_name = str(config.get("store_name") or "")
        for output_index, output_row in enumerate(output_rows, start=2):
            record = done.get(_key(output_row), {})
            note = str(record.get("note") or "")
            values = [
                store_name,
                output_row.order_number,
                output_row.recipient,
                str(config.get("phone") or ""),
                str(config.get("address1") or ""),
                str(config.get("address2") or ""),
                str(config.get("city") or ""),
                str(config.get("state") or ""),
                str(config.get("zip_code") or ""),
                f"https://www.amazon.com/dp/{output_row.sku}",
                record.get("price"),
                output_row.quantity,
                output_row.spec,
                note,
            ]
            for col, value in enumerate(values, start=1):
                ws.cell(row=output_index, column=col).value = value
        state.update({
            "state_version": STATE_VERSION,
            "input_excel": str(input_path.resolve()),
            "input_fingerprint": input_fingerprint,
            "total": total,
            "done": done,
            "status": "cancelled" if cancel_event is not None and cancel_event.is_set() else "completed",
            "updated_at": datetime.now().isoformat(),
        })
        _save_state(state_path, state)
        wb.save(target)
    finally:
        wb.close()
        if client is not None and hasattr(client, "close"):
            client.close()
    return target

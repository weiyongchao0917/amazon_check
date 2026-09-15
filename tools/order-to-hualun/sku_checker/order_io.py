"""Fixed-header input and Hualun output row mapping."""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

REQUIRED_HEADERS = ("订单编号", "平台SKU", "产品数量", "产品规格(原文)")


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip().lstrip("\t").strip()


def _quantity(value: Any) -> int | float:
    if isinstance(value, bool):
        raise ValueError("产品数量必须是数字")
    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = _clean(value).replace(",", "")
        try:
            number = float(text)
        except ValueError as exc:
            raise ValueError(f"产品数量不是有效数字: {text}") from exc
    return int(number) if number.is_integer() else number


@dataclass(frozen=True)
class OrderRow:
    row_num: int
    order_id: str
    platform_sku: str
    quantity: int | float
    spec: str


@dataclass(frozen=True)
class OutputRow:
    source: OrderRow
    order_number: str
    recipient: str

    @property
    def sku(self) -> str:
        return self.source.platform_sku

    @property
    def quantity(self) -> int | float:
        return self.source.quantity

    @property
    def spec(self) -> str:
        return self.source.spec


def load_order_rows(path: str | Path) -> list[OrderRow]:
    target = Path(path)
    if not target.is_file():
        raise ValueError(f"订单 Excel 不存在: {target}")
    wb = load_workbook(target, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            iterator = ws.iter_rows(values_only=True)
            try:
                raw_headers = next(iterator)
            except StopIteration:
                continue
            headers = [_clean(value) for value in raw_headers]
            if not all(header in headers for header in REQUIRED_HEADERS):
                continue
            indexes = {header: headers.index(header) for header in REQUIRED_HEADERS}
            result: list[OrderRow] = []
            for row_num, row in enumerate(iterator, start=2):
                if not row or all(_clean(value) == "" for value in row):
                    continue
                order_id = _clean(row[indexes["订单编号"]])
                sku = _clean(row[indexes["平台SKU"]])
                quantity = _quantity(row[indexes["产品数量"]])
                spec = _clean(row[indexes["产品规格(原文)"]])
                if not order_id or not sku:
                    raise ValueError(f"第 {row_num} 行缺少订单编号或平台SKU")
                result.append(OrderRow(row_num, order_id, sku, quantity, spec))
            return result
    finally:
        wb.close()
    raise ValueError("订单 Excel 没有找到固定表头：订单编号、平台SKU、产品数量、产品规格(原文)")


def make_output_rows(rows: list[OrderRow], recipient_prefix: str = "Claire-wyc") -> list[OutputRow]:
    totals: dict[str, int] = {}
    counters: dict[str, int] = {}
    for row in rows:
        totals[row.order_id] = totals.get(row.order_id, 0) + 1
    output: list[OutputRow] = []
    for row in rows:
        counters[row.order_id] = counters.get(row.order_id, 0) + 1
        number = row.order_id if totals[row.order_id] == 1 else f"{row.order_id}-{counters[row.order_id]}"
        output.append(OutputRow(row, number, f"{recipient_prefix}{row.order_id[-5:]}"))
    return output

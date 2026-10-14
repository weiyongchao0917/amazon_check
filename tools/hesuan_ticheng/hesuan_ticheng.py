"""订单利润核算程序。"""

from __future__ import annotations

import argparse
import math
import re
import sys
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Optional

import pandas as pd


EXCLUDED_STATUS_WORDS = ("退款", "拒绝", "驳回")


@dataclass
class CalculationResult:
    rows: list[dict[str, Any]]
    summary: dict[str, Any]


@dataclass
class InputFiles:
    main: Path
    purchase: Path
    product: Optional[Path] = None


def normalize_text(value: Any) -> str:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return ""
    return " ".join(str(value).replace("\ufeff", "").replace("\u200b", "").strip().split())


def normalize_id(value: Any) -> str:
    return unicodedata.normalize("NFKC", normalize_text(value))


def parse_number(value: Any) -> Optional[float]:
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = normalize_text(value).replace(",", "").replace("$", "")
    if not text or text.lower() in {"-", "--", "null", "none", "nan", "n/a"}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def base_order_number(value: Any) -> str:
    text = normalize_id(value)
    return re.sub(r"-\d+$", "", text) if text else ""


def is_excluded_purchase_status(value: Any) -> bool:
    status = normalize_text(value)
    return any(word in status for word in EXCLUDED_STATUS_WORDS)


def _find_column(frame: pd.DataFrame, aliases: Iterable[str], position: Optional[int] = None) -> Optional[str]:
    columns = {normalize_text(col): col for col in frame.columns}
    aliases = tuple(normalize_text(alias) for alias in aliases)
    for alias in aliases:
        if alias in columns:
            return columns[alias]
    for column in frame.columns:
        normalized = normalize_text(column)
        if any(alias and alias in normalized for alias in aliases):
            return column
    if position is not None and position < len(frame.columns):
        return frame.columns[position]
    return None


def _value(row: pd.Series, column: Optional[str]) -> Any:
    return row[column] if column is not None and column in row.index else None


def _number(row: pd.Series, column: Optional[str]) -> Optional[float]:
    return parse_number(_value(row, column))


def _amount(value: Optional[float]) -> Optional[float]:
    return None if value is None else round(float(value), 2)


def _read_csv(path: Path, nrows: Optional[int] = None) -> pd.DataFrame:
    last_error: Optional[Exception] = None
    for encoding in ("utf-8-sig", "gb18030", "utf-8"):
        try:
            return pd.read_csv(path, encoding=encoding, dtype=object, nrows=nrows)
        except (UnicodeDecodeError, pd.errors.ParserError) as exc:
            last_error = exc
    raise ValueError(f"无法读取 CSV 文件 {path}: {last_error}")


def read_input_table(path: Path) -> pd.DataFrame:
    if path.suffix.lower() in {".xlsx", ".xls", ".xlsm"}:
        return pd.read_excel(path, dtype=object)
    if path.suffix.lower() in {".csv", ".tsv"}:
        return _read_csv(path)
    raise ValueError(f"不支持的输入格式: {path}")


def find_input_files(folder: Path) -> InputFiles:
    folder = Path(folder)
    if not folder.is_dir():
        raise FileNotFoundError(f"工作目录不存在：{folder}")
    candidates: dict[str, list[Path]] = {"主表": [], "采购表": [], "商品表": []}
    errors = []
    for path in sorted(folder.iterdir()):
        if not path.is_file() or path.name.startswith("~$") or path.suffix.lower() not in {".xlsx", ".xlsm", ".csv"}:
            continue
        try:
            frame = _read_csv(path, nrows=0) if path.suffix.lower() == ".csv" else pd.read_excel(path, nrows=0)
        except Exception as exc:
            errors.append(f"{path.name}：{exc}")
            continue
        columns = {unicodedata.normalize("NFKC", normalize_text(col)).casefold() for col in frame.columns}

        def has(*aliases):
            return any(alias.casefold() in columns for alias in aliases)

        # 输出文件也含回款列，必须排除，避免被当成新的原始主表。
        if has("订单利润", "采用成本", "成本来源", "核对状态"):
            continue
        order = has("订单编号", "买家订单号", "订单号")
        if order and has("平台回款金额", "回款金额", "平台回款"):
            candidates["主表"].append(path)
        elif order and has("采购总金额", "采购金额", "采购成本", "商品累计金额") and has("采购状态", "订单状态", "状态"):
            candidates["采购表"].append(path)
        elif has("产品id", "商品id", "产品编号") and has("产品成本($)", "商品成本($)", "产品成本", "商品成本", "成本"):
            candidates["商品表"].append(path)
    missing = [kind for kind in ("主表", "采购表") if not candidates[kind]]
    if missing:
        message = f"目录中未识别到{'、'.join(missing)}。请选择包含原始 Excel/CSV 的目录；主表需要订单编号和平台回款金额，采购表需要买家订单号、采购总金额和订单状态。"
        if errors:
            message += "\n无法读取的文件：\n" + "\n".join(errors)
        raise FileNotFoundError(message)

    def latest(kind):
        return max(candidates[kind], key=lambda path: (path.stat().st_mtime, "副本" not in path.stem, path.name)) if candidates[kind] else None

    return InputFiles(latest("主表"), latest("采购表"), latest("商品表"))


def _main_columns(frame: pd.DataFrame) -> dict[str, Optional[str]]:
    return {
        "order_total": _find_column(frame, ("订单总金额", "订单金额"), 0),
        "platform_return": _find_column(frame, ("平台回款金额", "回款金额", "平台回款"), 1),
        "order_time": _find_column(frame, ("下单时间", "订单时间"), 2),
        "order_id": _find_column(frame, ("订单编号", "买家订单号", "订单号"), 4),
        "store": _find_column(frame, ("店铺名称", "店铺"), 5),
        "status": _find_column(frame, ("订单状态", "状态"), 6),
        "product_id": _find_column(frame, ("产品ID", "商品ID", "产品id", "商品id", "产品编号"), 7),
        "original_name": _find_column(frame, ("产品描述(原始)", "商品名称", "产品名称"), 8),
        "chinese_name": _find_column(frame, ("产品描述(中文)", "中文商品名称"), 9),
        "quantity": _find_column(frame, ("产品数量", "商品数量", "采购数量", "数量"), 10),
    }


def _purchase_columns(frame: pd.DataFrame) -> dict[str, Optional[str]]:
    return {
        "order_id": _find_column(frame, ("买家订单号", "订单编号", "订单号"), 2),
        "quantity": _find_column(frame, ("采购数量", "商品数量", "数量"), 14),
        "cost": _find_column(frame, ("采购总金额", "采购金额", "采购成本", "商品累计金额"), 23),
        "status": _find_column(frame, ("采购状态", "订单状态", "状态"), 24),
    }


def _product_columns(frame: pd.DataFrame) -> dict[str, Optional[str]]:
    return {
        "product_id": _find_column(frame, ("产品id", "产品ID", "商品id", "商品ID", "产品编号"), 0),
        "unit_cost": _find_column(frame, ("产品成本($)", "商品成本($)", "产品成本", "商品成本", "成本"), 2),
    }


def _build_product_costs(product_df: Optional[pd.DataFrame]) -> dict[str, float]:
    if product_df is None or product_df.empty:
        return {}
    columns = _product_columns(product_df)
    if columns["product_id"] is None or columns["unit_cost"] is None:
        return {}
    result: dict[str, float] = {}
    for _, row in product_df.iterrows():
        product_id = normalize_id(_value(row, columns["product_id"]))
        unit_cost = parse_number(_value(row, columns["unit_cost"]))
        if not product_id or unit_cost is None:
            continue
        if product_id in result and not math.isclose(result[product_id], unit_cost, abs_tol=0.000001):
            raise ValueError(f"商品表产品 ID {product_id} 存在不同成本，无法安全计算")
        result[product_id] = unit_cost
    return result


def calculate_profit(main_df: pd.DataFrame, purchase_df: pd.DataFrame, product_df: Optional[pd.DataFrame] = None, logistics_fee: float = 1.0) -> CalculationResult:
    main_columns = _main_columns(main_df)
    purchase_columns = _purchase_columns(purchase_df)
    product_costs = _build_product_costs(product_df)
    purchase_groups: dict[str, list[dict[str, Any]]] = {}
    if purchase_columns["order_id"] is not None:
        for source_index, row in purchase_df.iterrows():
            suborder = normalize_id(_value(row, purchase_columns["order_id"]))
            base = base_order_number(suborder)
            if not base:
                continue
            purchase_groups.setdefault(base, []).append({
                "source_index": int(source_index) + 2,
                "suborder": suborder,
                "cost": _number(row, purchase_columns["cost"]),
                "quantity": _number(row, purchase_columns["quantity"]),
                "status": normalize_text(_value(row, purchase_columns["status"])),
                "excluded": is_excluded_purchase_status(_value(row, purchase_columns["status"])),
            })

    rows: list[dict[str, Any]] = []
    for sequence, (_, row) in enumerate(main_df.iterrows(), 1):
        raw_order_id = _value(row, main_columns["order_id"])
        order_id = normalize_id(raw_order_id)
        base = base_order_number(order_id)
        revenue = _number(row, main_columns["platform_return"])
        if revenue is None:
            revenue = _number(row, main_columns["order_total"])
        quantity = _number(row, main_columns["quantity"])
        quantity = 1 if quantity is None else quantity
        product_id = normalize_id(_value(row, main_columns["product_id"]))
        matches = purchase_groups.get(base, []) if base else []
        included = [item for item in matches if not item["excluded"]]
        excluded = [item for item in matches if item["excluded"]]
        costs = [item["cost"] for item in included if item["cost"] is not None]
        purchase_cost = _amount(sum(costs)) if costs and len(costs) == len(included) else None
        excluded_amount = _amount(sum(item["cost"] for item in excluded if item["cost"] is not None))
        product_unit_cost = product_costs.get(product_id)
        product_total = _amount(product_unit_cost * quantity) if product_unit_cost is not None else None
        if not order_id:
            status = "订单编号为空，待核对"
        elif purchase_cost is not None:
            status = "已匹配采购订单"
        elif product_total is not None:
            status = "已匹配商品成本"
        elif excluded and not included:
            status = "仅有退款/拒绝采购记录，待商品表"
        elif matches:
            status = "已找到采购订单但成本缺失"
        else:
            status = "未找到采购订单，待商品表"
        if purchase_cost is not None:
            adopted_cost, cost_source = purchase_cost, "采购订单"
            logistics = _amount(len(included) * logistics_fee) or 0.0
        elif product_total is not None:
            adopted_cost, cost_source = product_total, "商品表"
            logistics = _amount(logistics_fee) or 0.0
        else:
            adopted_cost, cost_source, logistics = None, "", 0.0
        profit = _amount(revenue - adopted_cost - logistics) if revenue is not None and adopted_cost is not None else None
        rows.append({
            "序号": sequence, "订单编号": raw_order_id, "订单编号(归并)": base,
            "店铺名称": _value(row, main_columns["store"]), "订单状态": _value(row, main_columns["status"]),
            "下单时间": _value(row, main_columns["order_time"]), "平台回款金额": _amount(revenue),
            "商品ID": _value(row, main_columns["product_id"]), "商品数量": quantity,
            "商品名称": _value(row, main_columns["chinese_name"]) or _value(row, main_columns["original_name"]),
            "计入采购行数": len(included), "排除采购行数": len(excluded),
            "计入采购子订单": ", ".join(item["suborder"] for item in included),
            "排除采购子订单": ", ".join(item["suborder"] for item in excluded),
            "采购商品数量合计": _amount(sum(item["quantity"] for item in included if item["quantity"] is not None)),
            "采购成本": purchase_cost, "排除采购金额": excluded_amount,
            "商品表匹配行数": 1 if product_total is not None else 0, "商品成本单价": _amount(product_unit_cost),
            "商品成本合计": product_total, "采用成本": adopted_cost, "物流成本": logistics,
            "订单利润": profit, "成本来源": cost_source, "核对状态": status,
        })
    profit_values = [row["订单利润"] for row in rows if row["订单利润"] is not None]
    loss_details = [
        {"订单编号": row["订单编号"], "亏损金额": _amount(abs(row["订单利润"]))}
        for row in rows
        if row["订单利润"] is not None and row["订单利润"] < 0
    ]
    summary = {
        "主订单数": len(rows), "已计算订单数": len(profit_values), "待核对订单数": len(rows) - len(profit_values),
        "亏损订单数": len(loss_details), "亏损订单明细": loss_details,
        "利润合计": _amount(sum(profit_values)), "采购订单成本合计": _amount(sum(row["采购成本"] or 0 for row in rows)),
        "商品表成本合计": _amount(sum(row["商品成本合计"] or 0 for row in rows)),
        "物流成本合计": _amount(sum(row["物流成本"] or 0 for row in rows)),
        "排除采购行数": sum(row["排除采购行数"] for row in rows),
        "排除采购金额": _amount(sum(row["排除采购金额"] or 0 for row in rows)),
    }
    return CalculationResult(rows, summary)


def export_result(result: CalculationResult, output_path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "汇总"
    summary_sheet["A1"] = "订单利润核算"
    summary_sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    summary_sheet["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    summary_sheet.merge_cells("A1:B1")
    summary_sheet["A2"] = "计算规则"
    summary_sheet["B2"] = "有效采购子订单各计 $1 物流；无有效采购记录时，商品表回退按一个履约单计 $1。"
    summary_sheet.merge_cells("B2:E2")
    metric_values = ((label, value) for label, value in result.summary.items() if label not in {"亏损订单数", "亏损订单明细"})
    for row_index, (label, value) in enumerate(metric_values, 4):
        summary_sheet.cell(row_index, 1, label)
        summary_sheet.cell(row_index, 2, value)
    summary_sheet.column_dimensions["A"].width = 24
    summary_sheet.column_dimensions["B"].width = 80
    headers = [
        "序号", "订单编号", "订单编号(归并)", "店铺名称", "订单状态", "平台回款金额",
        "采用成本", "物流成本", "订单利润", "商品ID", "商品数量", "商品名称",
        "计入采购行数", "排除采购行数", "计入采购子订单", "排除采购子订单",
        "采购商品数量合计", "采购成本", "排除采购金额", "商品表匹配行数", "商品成本单价",
        "商品成本合计", "成本来源", "核对状态",
    ]

    def write_detail_sheet(sheet, detail_rows):
        for col_index, header in enumerate(headers, 1):
            cell = sheet.cell(1, col_index, header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row_index, row in enumerate(detail_rows, 2):
            for col_index, header in enumerate(headers, 1):
                cell = sheet.cell(row_index, col_index, row.get(header))
                if header in {"订单编号", "订单编号(归并)", "商品ID", "计入采购子订单", "排除采购子订单"}:
                    cell.number_format = "@"
                if header in {"平台回款金额", "采购成本", "排除采购金额", "商品成本单价", "商品成本合计", "采用成本", "物流成本", "订单利润"}:
                    cell.number_format = '$#,##0.00;[Red]-$#,##0.00'
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(detail_rows) + 1)}"
        sheet.row_dimensions[1].height = 32
        for col_index, header in enumerate(headers, 1):
            width = min(max(12, len(str(header)) * 2), 28)
            if header in {"计入采购子订单", "排除采购子订单", "核对状态"}:
                width = 34
            sheet.column_dimensions[get_column_letter(col_index)].width = width
        status_col = get_column_letter(headers.index("核对状态") + 1)
        profit_col = get_column_letter(headers.index("订单利润") + 1)
        last_row = max(2, len(detail_rows) + 1)
        sheet.conditional_formatting.add(f"{status_col}2:{status_col}{last_row}", FormulaRule(formula=[f'ISNUMBER(SEARCH("待",{status_col}2))'], fill=PatternFill("solid", fgColor="FFF2CC")))
        sheet.conditional_formatting.add(f"{profit_col}2:{profit_col}{last_row}", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor="F4CCCC")))

    profit_rows = [row for row in result.rows if row.get("订单利润") is not None and row["订单利润"] >= 0]
    loss_rows = [row for row in result.rows if row.get("订单利润") is not None and row["订单利润"] < 0]
    pending_rows = [row for row in result.rows if row.get("订单利润") is None]
    for sheet_name, detail_rows in (("盈利订单", profit_rows), ("亏损订单", loss_rows), ("待核对订单", pending_rows)):
        write_detail_sheet(workbook.create_sheet(sheet_name), detail_rows)
    workbook.save(output_path)
    return output_path


def main(argv: Optional[list[str]] = None) -> int:
    parser = argparse.ArgumentParser(description="根据订单、采购订单和商品表核算每个订单利润")
    parser.add_argument("--folder", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument("--output", type=Path, default=None)
    parser.add_argument("--logistics-fee", type=float, default=1.0)
    args = parser.parse_args(argv)
    try:
        files = find_input_files(args.folder)
        result = calculate_profit(read_input_table(files.main), read_input_table(files.purchase), read_input_table(files.product) if files.product else None, args.logistics_fee)
        output_path = args.output or args.folder / "输出" / f"订单利润核算-{datetime.now():%Y%m%d-%H%M%S}.xlsx"
        export_result(result, output_path)
        print(f"主表: {files.main.name}\n采购表: {files.purchase.name}\n商品表: {files.product.name if files.product else '未提供'}\n输出: {output_path}")
        for key, value in result.summary.items():
            print(f"{key}: {value}")
        return 0
    except Exception as exc:
        print(f"核算失败: {exc}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())

import pandas as pd
from hesuan_ticheng import calculate_profit, export_result, find_input_files


def test_detects_arbitrary_filenames_by_headers(tmp_path):
    pd.DataFrame([main_row()]).to_csv(tmp_path / "本周回款.csv", index=False)
    pd.DataFrame([purchase_row("100", 20)]).to_csv(tmp_path / "供应商记录.csv", index=False)
    pd.DataFrame([{"产品id": "P1", "产品成本($)": 4}]).to_csv(tmp_path / "成本资料.csv", index=False)
    files = find_input_files(tmp_path)
    assert files.main.name == "本周回款.csv"
    assert files.purchase.name == "供应商记录.csv"
    assert files.product.name == "成本资料.csv"


def test_detection_ignores_output_and_unrelated_tables(tmp_path):
    import pytest
    pd.DataFrame([main_row()]).to_csv(tmp_path / "回款.csv", index=False)
    pd.DataFrame([{"名称": "不是采购表"}]).to_csv(tmp_path / "采购订单.csv", index=False)
    with pytest.raises(FileNotFoundError, match="采购表"):
        find_input_files(tmp_path)
    pd.DataFrame([purchase_row("100", 20)]).to_csv(tmp_path / "真实采购.csv", index=False)
    output = dict(main_row(), **{"采用成本": 20, "订单利润": 29})
    pd.DataFrame([output]).to_csv(tmp_path / "核算结果.csv", index=False)
    files = find_input_files(tmp_path)
    assert files.main.name == "回款.csv"
    assert files.product is None


def test_detection_uses_newest_table_and_skips_locked_or_broken_files(tmp_path):
    import os
    for name in ("较旧回款.csv", "较新回款.csv"):
        pd.DataFrame([main_row()]).to_csv(tmp_path / name, index=False, encoding="gb18030")
    os.utime(tmp_path / "较旧回款.csv", (1000, 1000))
    pd.DataFrame([purchase_row("100", 20)]).to_csv(tmp_path / "采购.csv", index=False)
    (tmp_path / "~$正在编辑.xlsx").write_text("lock")
    (tmp_path / "损坏.xlsx").write_text("invalid")
    files = find_input_files(tmp_path)
    assert files.main.name == "较新回款.csv"


def main_row(order="100", revenue=50, product_id="P1", quantity=2):
    return {
        "订单总金额": revenue,
        "平台回款金额": revenue,
        "下单时间": "2026-10-13 10:00:00",
        "订单编号": order,
        "店铺名称": "shop",
        "订单状态": "已完成",
        "产品ID": product_id,
        "产品数量": quantity,
        "产品描述(原始)": "item",
        "产品描述(中文)": "商品",
    }


def purchase_row(order, cost, status="已收货", quantity=1):
    return {"买家订单号": order, "采购总金额": cost, "采购状态": status, "采购数量": quantity}


def test_suborders_are_merged_and_each_suborder_costs_one_dollar():
    main = pd.DataFrame([main_row()])
    purchase = pd.DataFrame([
        purchase_row("100-1", 10),
        purchase_row("100-2", 12),
    ])
    result = calculate_profit(main, purchase, None)
    row = result.rows[0]
    assert row["采购成本"] == 22
    assert row["物流成本"] == 2
    assert row["订单利润"] == 26


def test_product_cost_fallback_uses_unit_cost_times_quantity_and_one_dollar_logistics():
    main = pd.DataFrame([main_row(order="200", revenue=30, product_id="P2", quantity=3)])
    purchase = pd.DataFrame(columns=["买家订单号", "采购总金额", "采购状态", "采购数量"])
    products = pd.DataFrame([{"产品id": "P2", "产品成本($)": 4}])
    result = calculate_profit(main, purchase, products)
    row = result.rows[0]
    assert row["商品成本合计"] == 12
    assert row["物流成本"] == 1
    assert row["订单利润"] == 17
    assert row["核对状态"] == "已匹配商品成本"


def test_refunded_rejected_and_rejected_orders_are_excluded():
    main = pd.DataFrame([main_row(order="300", revenue=40)])
    purchase = pd.DataFrame([
        purchase_row("300-1", 9, status="已退款"),
        purchase_row("300-2", 8, status="已拒绝"),
        purchase_row("300-3", 7, status="已驳回订单"),
    ])
    result = calculate_profit(main, purchase, None)
    row = result.rows[0]
    assert row["采购成本"] is None
    assert row["排除采购行数"] == 3
    assert row["排除采购金额"] == 24
    assert row["核对状态"] == "仅有退款/拒绝采购记录，待商品表"


def test_valid_purchase_cost_has_priority_over_product_cost():
    main = pd.DataFrame([main_row(order="400", revenue=50, product_id="P4", quantity=2)])
    purchase = pd.DataFrame([purchase_row("400", 20)])
    products = pd.DataFrame([{"产品id": "P4", "产品成本($)": 3}])
    result = calculate_profit(main, purchase, products)
    row = result.rows[0]
    assert row["采用成本"] == 20
    assert row["成本来源"] == "采购订单"
    assert row["订单利润"] == 29


def test_order_contains_total_amount_and_fixed_expense():
    main = pd.DataFrame([main_row(order="450", revenue=40)])
    main.loc[0, "订单总金额"] = 50
    purchase = pd.DataFrame([purchase_row("450", 20)])
    result = calculate_profit(main, purchase, None)
    row = result.rows[0]
    assert row["订单总金额"] == 50
    assert row["固定支出"] == 10
    assert row["订单利润"] == 19


def test_product_group_combines_multiple_links_and_calculates_unit_fixed_expense():
    main = pd.DataFrame([
        main_row(order="501", revenue=40, product_id="L1", quantity=2),
        main_row(order="502", revenue=30, product_id="L2", quantity=1),
    ])
    main.loc[0, "订单总金额"] = 50
    main.loc[1, "订单总金额"] = 40
    purchase = pd.DataFrame([purchase_row("501", 10), purchase_row("502", 8)])
    products = pd.DataFrame([
        {"商品组ID": "PG-000001", "产品id": "L1", "产品成本($)": 4},
        {"商品组ID": "PG-000001", "产品id": "L2", "产品成本($)": 5},
    ])
    result = calculate_profit(main, purchase, products)
    assert result.rows[0]["商品组ID"] == "PG-000001"
    assert len(result.product_rows) == 1
    summary = result.product_rows[0]
    assert summary["商品组ID"] == "PG-000001"
    assert summary["销售数量"] == 3
    assert summary["固定支出合计"] == 20
    assert summary["单件固定支出"] == 6.67
    assert summary["产品ID数量"] == 2


def test_missing_product_group_does_not_merge_different_products():
    main = pd.DataFrame([
        main_row(order="503", revenue=20, product_id="L3"),
        main_row(order="504", revenue=20, product_id="L4"),
    ])
    purchase = pd.DataFrame([purchase_row("503", 5), purchase_row("504", 6)])
    products = pd.DataFrame([
        {"产品id": "L3", "产品成本($)": 3},
        {"产品id": "L4", "产品成本($)": 4},
    ])
    result = calculate_profit(main, purchase, products)
    assert {row["商品组ID"] for row in result.product_rows} == {"未归组-L3", "未归组-L4"}


def test_blank_order_id_is_not_matched():
    main = pd.DataFrame([main_row(order="", revenue=20)])
    purchase = pd.DataFrame([purchase_row("", 5)])
    result = calculate_profit(main, purchase, None)
    assert result.rows[0]["核对状态"] == "订单编号为空，待核对"


def test_loss_orders_are_listed_with_negative_profit():
    main = pd.DataFrame([main_row(order="500", revenue=10)])
    purchase = pd.DataFrame([purchase_row("500", 20)])
    result = calculate_profit(main, purchase, None)
    assert result.rows[0]["订单利润"] == -11
    assert result.summary["亏损订单明细"] == [{"订单编号": "500", "亏损金额": 11}]


def test_export_splits_profit_and_loss_sheets_and_orders_profit_columns(tmp_path):
    main = pd.DataFrame([
        main_row(order="600", revenue=50),
        main_row(order="601", revenue=10),
    ])
    purchase = pd.DataFrame([
        purchase_row("600", 20),
        purchase_row("601", 20),
    ])
    result = calculate_profit(main, purchase, None)
    output = export_result(result, tmp_path / "result.xlsx")

    import openpyxl
    workbook = openpyxl.load_workbook(output, data_only=True)
    assert workbook.sheetnames == ["汇总", "盈利订单", "亏损订单", "待核对订单", "商品成本汇总"]
    assert workbook["汇总"]["D3"].value is None
    profit_headers = [cell.value for cell in workbook["盈利订单"][1]]
    loss_headers = [cell.value for cell in workbook["亏损订单"][1]]
    assert "下单时间" not in profit_headers
    assert profit_headers[profit_headers.index("平台回款金额") + 1:profit_headers.index("平台回款金额") + 4] == ["固定支出", "采用成本", "物流成本"]
    assert loss_headers == profit_headers
    assert workbook["亏损订单"].cell(2, loss_headers.index("订单利润") + 1).value == -11
    assert workbook["商品成本汇总"].max_row == 2
